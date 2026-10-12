"""Реестр писем: из letters_raw.json (разбор писем) + исходные файлы -> pack/Письма/*.pdf и лист «Письма» в Требования_ОКК.xlsx.
  python3 letters/build_letters.py <папка исходников писем> <папка pack>
"""
import json, os, re, sys, shutil, subprocess, tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
SRC, PACK = sys.argv[1], sys.argv[2]
L = json.load(open(os.path.join(HERE, 'letters_raw.json')))
OVR = json.load(open(os.path.join(HERE, 'overrides.json'))) if os.path.exists(os.path.join(HERE, 'overrides.json')) else {}
L = [x for x in L if x['folder'] not in OVR.get('drop', [])]
for x in L:
    x.update(OVR.get('patch', {}).get(x['folder'], {}))

dn = lambda d: tuple(reversed([int(p) for p in d.split('.')])) if re.match(r'\d\d\.\d\d\.\d{4}$', d or '') else (0, 0, 0)
L.sort(key=lambda x: dn(x['date']), reverse=True)
key = lambda n: re.sub(r'[^0-9A-Za-zА-Яа-я]', '', n or '').upper()
tail = lambda n: re.sub(r'^.*-', '', n or '')
by = {key(x['number']): x for x in L}
def find(n):
    k = key(n)
    if k in by: return by[k]
    for x in L:  # «3024-10-226-0392» и т.п.: по хвосту номера и серии
        if tail(x['number']) == tail(n) and x['number'].split('-')[1:2] == n.split('-')[1:2]: return x
    return None
lbl = lambda x: f"№ {x['number']} от {x['date']}"

cancelled, changed = {}, {}
for x in L:
    for c in x['cancels']:
        t = find(c['number'])
        if t and t is not x: cancelled.setdefault(t['number'], []).append(x)
    for c in x['changes']:
        t = find(c['number'])
        if t and t is not x: changed.setdefault(t['number'], []).append((x, c.get('what', '')))

# файлы
docs_dir = os.path.join(PACK, 'Документы'); out_dir = os.path.join(PACK, 'Письма')
shutil.rmtree(out_dir, ignore_errors=True); os.makedirs(out_dir)
have = {f.lower() for f in os.listdir(docs_dir)}
def to_pdf(src, dst):
    if src.lower().endswith(('.jpg', '.jpeg', '.png')):
        from PIL import Image
        im = Image.open(src).convert('RGB'); im.thumbnail((1800, 1800)); im.save(dst, 'PDF', resolution=150); return
    tmp = dst + '.gs.pdf'
    r = subprocess.run(['gs', '-q', '-sDEVICE=pdfwrite', '-dPDFSETTINGS=/ebook', '-dCompatibilityLevel=1.5', '-dNOPAUSE', '-dBATCH',
                        '-dDetectDuplicateImages=true', f'-sOutputFile={tmp}', src], capture_output=True)
    if r.returncode == 0 and os.path.exists(tmp) and 0 < os.path.getsize(tmp) < os.path.getsize(src) * 0.9:
        os.replace(tmp, dst)
    else:
        if os.path.exists(tmp): os.remove(tmp)
        shutil.copy(src, dst)
safe = lambda s: re.sub(r'[\\/:*?"<>|]', '-', s)
for x in L:
    folder = os.path.join(SRC, x['folder'])
    base = safe(f"Письмо {x['number']} от {x['date']}")
    x['file'] = ''
    if x.get('main_file') and os.path.exists(os.path.join(folder, x['main_file'])):
        fn = base + '.pdf'
        if fn.lower() in have: x['file'] = fn  # уже есть среди документов требований
        else: to_pdf(os.path.join(folder, x['main_file']), os.path.join(out_dir, fn)); x['file'] = fn
    att = []
    i = 0
    for a in x['attachments']:
        p = os.path.join(folder, a['file'])
        if not os.path.exists(p) or a['file'] == x.get('main_file'): continue
        if a['file'].lower().endswith(('.pdf', '.jpg', '.jpeg', '.png')):
            if a['file'].lower().endswith('.pdf') and any(b['file'].lower().endswith('.docx') for b in []): pass
            i += 1; fn = f"{base} — приложение {i}.pdf"
            to_pdf(p, os.path.join(out_dir, fn)); att.append((fn, a['what']))
        elif a['file'].lower().endswith('.xlsx'):
            att.append(('', a['what'] + ' (таблица Excel — в приложении не открывается)'))
    x['att'] = att

# Excel
xp = os.path.join(PACK, 'Требования_ОКК.xlsx')
wb = openpyxl.load_workbook(xp)
if 'Письма' in wb.sheetnames: del wb['Письма']
ws = wb.create_sheet('Письма', index=wb.sheetnames.index('Документы') + 1)
HEAD = ['Номер', 'Дата', 'От кого', 'Кому', 'Ответ на', 'Тема', 'Суть', 'Ключевые слова', 'Категория', 'Конструктив',
        'Статус', 'Отменено письмом', 'Отменяет', 'Изменено письмом', 'Изменяет', 'Примечание', 'Файл PDF', 'Приложения']
ws.append(HEAD)
for x in L:
    cb = cancelled.get(x['number'], [])
    ch = changed.get(x['number'], [])
    status = 'Отменено' if cb else ('Изменено частично' if ch else 'Действует')
    cancels = '; '.join((lbl(find(c['number'])) if find(c['number']) else '№ ' + c['number']) for c in x['cancels'])
    changes = '\n'.join(((lbl(find(c['number'])) if find(c['number']) else '№ ' + c['number']) + (': ' + c['what'] if c.get('what') else '')) for c in x['changes'])
    note = '; '.join(filter(None, [('Действует только: ' + x['temporary']) if x.get('temporary') else '', x.get('note', '')]))
    ws.append([x['number'], x['date'], x['from'], x['to'], x['reply_to'], x['subject'], x['summary'], ', '.join(x['keywords']),
               x['category'], '; '.join(x['kons']), status, '; '.join(lbl(c) for c in cb), cancels,
               '\n'.join(lbl(c) + (': ' + w if w else '') for c, w in ch), changes, note, x['file'],
               '\n'.join((f or '—') + ' | ' + w for f, w in x['att'])])
W = [20, 11, 18, 22, 18, 40, 70, 40, 16, 26, 14, 26, 26, 34, 34, 34, 34, 50]
for i, w in enumerate(W): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1D3A5C'); c.alignment = Alignment(wrap_text=True, vertical='center')
for r in ws.iter_rows(min_row=2):
    for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
    if r[10].value == 'Отменено':
        for c in r: c.fill = PatternFill('solid', fgColor='FBE5E5')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
# у документов требований: подставить PDF, если его не было (напр. 4335)
wd = wb['Документы']
for r in wd.iter_rows(min_row=2):
    code = str(r[0].value or '')
    if code.startswith('Письмо') and not r[4].value:
        x = find(code.replace('Письмо', '').strip())
        if x and x['file']:
            src = os.path.join(out_dir, x['file'])
            if os.path.exists(src): shutil.copy(src, os.path.join(docs_dir, x['file'])); os.remove(src)
            r[4].value = x['file']
            if r[2].value and '(PDF нет)' in str(r[2].value): r[2].value = str(r[2].value).replace(' (PDF нет)', '')
            if r[5].value and str(r[5].value).startswith('PDF письма не загружен.'): r[5].value = str(r[5].value).replace('PDF письма не загружен. ', '')
wb.save(xp)
tot = sum(os.path.getsize(os.path.join(out_dir, f)) for f in os.listdir(out_dir))
print('писем', len(L), 'отменено', sum(1 for x in L if x['number'] in cancelled), 'изменено', len(changed), 'файлов', len(os.listdir(out_dir)), 'МБ', round(tot / 1e6, 1))
