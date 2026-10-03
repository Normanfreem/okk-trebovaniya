"""Сборка пакета данных ОКК из папки (Требования_ОКК.xlsx + Документы/*.pdf).

  python3 pack_tool.py <папка> <дата ДД.ММ.ГГГГ> [--pwd ПАРОЛЬ --out ПАПКА]

Проверяет Excel по правилам RULES.md, заполняет пустой «Вид материала»
по правилам cat_for(), ставит дату на лист «Инфо», собирает zip
и (если дан пароль) шифрует в data.okk + version.json через encrypt.py.
"""
import sys, os, re, zipfile, subprocess, argparse
import openpyxl

CAT_ORDER = ['Щебень', 'Песок и грунт', 'Монолитный бетон', 'Сборные ЖБ и бетонные изделия', 'Арматура',
             'Растворы', 'Гидроизоляция', 'Материалы для швов', 'Геосинтетика', 'Камень',
             'Металлоконструкции и окраска']
PRECAST = ('звень', 'стенки откосн', 'стенка откосн', 'фундаментные блоки', 'бетонные блоки', 'плит', 'оголов')


def cat_for(el, bl, mat):
    """Вид материала по умолчанию (если в Excel не заполнен)."""
    m = (mat or '').lower(); e = (el or '').lower()
    if bl == 'Приёмка': return ''
    if 'отклонен' in m or 'толщина защитного слоя' in m: return ''
    if 'щеб' in m: return 'Щебень'
    if 'арматур' in m: return 'Арматура'
    if 'досыпка' in m or 'песок' in m or 'грунт' in m and 'грунтовк' not in m: return 'Песок и грунт'
    if 'подливочн' in m or 'выравнивающ' in m or 'раствор' in m: return 'Растворы'
    if 'шв' in e and 'бетон' not in m and 'заполнение швов' not in m: return 'Материалы для швов'
    if 'геосинтет' in m or 'геотекст' in m: return 'Геосинтетика'
    if 'камен' in m or 'рисберма' in m: return 'Камень'
    if 'металл' in m or 'окраска' in m: return 'Металлоконструкции и окраска'
    if 'гидроизол' in e and 'защитный слой' not in m: return 'Гидроизоляция'
    if any(p in m for p in PRECAST): return 'Сборные ЖБ и бетонные изделия'
    return 'Монолитный бетон'


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('folder'); ap.add_argument('asof')
    ap.add_argument('--pwd'); ap.add_argument('--out', default='out'); ap.add_argument('--name', default='ОКК')
    a = ap.parse_args()
    xp = os.path.join(a.folder, 'Требования_ОКК.xlsx')
    wb = openpyxl.load_workbook(xp)
    wr, wd = wb['Требования'], wb['Документы']
    H = [str(c.value or '').strip().lower() for c in wr[1]]
    col = lambda *names: next(i for i, h in enumerate(H) if any(h.startswith(n) for n in names))
    ci = dict(id=col('id'), el=col('элемент'), bl=col('блок'), mat=col('материал'), cat=col('вид материала'),
              text=col('требование'), doc=col('документ'), page=col('стр'), st=col('статус'))
    docs = {str(r[0].value).strip(): str(r[4].value or '').strip() for r in wd.iter_rows(min_row=2) if r[0].value}
    errs, warn, ids = [], [], set()
    for r in wr.iter_rows(min_row=2):
        v = lambda k: '' if r[ci[k]].value is None else str(r[ci[k]].value).strip()
        if not v('id') and not v('text'): continue
        rid = v('id')
        if rid in ids: errs.append(f'{rid}: повтор ID')
        ids.add(rid)
        if v('bl') not in ('Материал', 'Устройство', 'Приёмка'): errs.append(f'{rid}: Блок «{v("bl")}»')
        dl = [x.strip() for x in v('doc').split(';') if x.strip()]
        pl = [x.strip() for x in v('page').split(';') if x.strip()]
        for d in dl:
            if d not in docs: errs.append(f'{rid}: документ «{d}» нет на листе «Документы»')
        if pl and len(pl) != len(dl): errs.append(f'{rid}: число страниц ≠ числу документов')
        if not v('cat') and v('bl') == 'Материал':
            c = cat_for(v('el'), v('bl'), v('mat'))
            if c: r[ci['cat']].value = c; warn.append(f'{rid}: вид материала → {c}')
        elif v('cat') and v('cat') not in CAT_ORDER: warn.append(f'{rid}: новый вид материала «{v("cat")}» (добавить в CAT_ORDER приложения)')
        if re.search(r'мелкозернист\w* B\d', v('text')): warn.append(f'{rid}: «мелкозернистый Bxx» — писать «Класс по прочности - Bxx»')
    for code, f in docs.items():
        if f and not os.path.exists(os.path.join(a.folder, 'Документы', f)): errs.append(f'Документ {code}: нет файла {f}')
    wi = wb['Инфо']
    for r in wi.iter_rows(min_row=1):
        if str(r[0].value).strip().lower().startswith('актуально'): r[1].value = a.asof
    print('\n'.join(['ОШИБКИ:'] + errs if errs else ['ошибок нет']))
    if warn: print('\n'.join(['Замечания:'] + warn))
    if errs: sys.exit(1)
    wb.save(xp)
    os.makedirs(a.out, exist_ok=True)
    d, m, y = a.asof.split('.')
    zp = os.path.join(a.out, f'{a.name}_{d}.{m}.{y[-2:]}.zip')
    with zipfile.ZipFile(zp, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(a.folder):
            for f in files:
                p = os.path.join(root, f); z.write(p, os.path.relpath(p, a.folder))
    print('zip:', zp, os.path.getsize(zp))
    if a.pwd:
        subprocess.check_call([sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'encrypt.py'), zp, a.pwd, a.out, a.asof])
        print('data.okk + version.json →', a.out)


if __name__ == '__main__':
    main()
