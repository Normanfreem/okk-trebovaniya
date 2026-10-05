"""Требования ОКК v2 — сервер: вход, права, пакеты данных, журнал, предложения.

Запуск: uvicorn app:app --host 127.0.0.1 --port 8001
Данные: каталог из переменной OKK_DATA (по умолчанию /opt/okk-data).
"""
import base64, csv, hashlib, hmac, io, json, os, secrets, sqlite3, threading, time, zipfile
from datetime import datetime, timezone

import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from fastapi import Body, Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles

DATA = os.environ.get('OKK_DATA', '/opt/okk-data')
WEB = os.environ.get('OKK_WEB', os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'web'))
DB_PATH = os.path.join(DATA, 'okk.sqlite')
os.makedirs(os.path.join(DATA, 'versions'), exist_ok=True)

_lock = threading.Lock()


def db():
    c = sqlite3.connect(DB_PATH, timeout=10)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA journal_mode=WAL')
    c.execute('PRAGMA foreign_keys=ON')
    return c


SCHEMA = '''
CREATE TABLE IF NOT EXISTS levels(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, perms TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, login TEXT UNIQUE NOT NULL, name TEXT NOT NULL DEFAULT '',
  pw_hash TEXT NOT NULL, pw_salt TEXT NOT NULL, level_id INTEGER REFERENCES levels(id), active INTEGER NOT NULL DEFAULT 1,
  is_admin INTEGER NOT NULL DEFAULT 0, offline_days INTEGER NOT NULL DEFAULT 35, dk TEXT NOT NULL,
  created TEXT NOT NULL, last_seen TEXT, note TEXT NOT NULL DEFAULT '');
CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  device TEXT, created TEXT NOT NULL, last_seen TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, ts TEXT NOT NULL, recv TEXT NOT NULL, user_id INTEGER, login TEXT,
  type TEXT NOT NULL, data TEXT, device TEXT, ip TEXT, offline INTEGER NOT NULL DEFAULT 0);
CREATE INDEX IF NOT EXISTS ev_ts ON events(ts);
CREATE INDEX IF NOT EXISTS ev_user ON events(user_id);
CREATE TABLE IF NOT EXISTS suggestions(id INTEGER PRIMARY KEY, ts TEXT NOT NULL, user_id INTEGER, login TEXT, kon TEXT, card TEXT,
  ids TEXT, text TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'new', note TEXT NOT NULL DEFAULT '', decided TEXT);
CREATE TABLE IF NOT EXISTS versions(id INTEGER PRIMARY KEY, ts TEXT NOT NULL, fname TEXT, asof TEXT, size INTEGER,
  kons TEXT, active INTEGER NOT NULL DEFAULT 0, by_login TEXT);
'''

DEFAULT_LEVELS = [
    ('Полный доступ', {'kons': ['*'], 'pdf': True, 'export': True, 'suggest': True}),
    ('Просмотр всего', {'kons': ['*'], 'pdf': True, 'export': True, 'suggest': False}),
    ('Ограниченный', {'kons': [], 'pdf': True, 'export': False, 'suggest': False}),
]


def init_db():
    with db() as c:
        c.executescript(SCHEMA)
        if not c.execute('SELECT 1 FROM levels').fetchone():
            for n, p in DEFAULT_LEVELS:
                c.execute('INSERT INTO levels(name,perms) VALUES(?,?)', (n, json.dumps(p, ensure_ascii=False)))


def now():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')


def hash_pw(pw, salt=None):
    salt = salt or secrets.token_hex(16)
    h = hashlib.scrypt(pw.encode(), salt=bytes.fromhex(salt), n=2 ** 14, r=8, p=1, dklen=32).hex()
    return h, salt


def check_pw(pw, h, salt):
    return hmac.compare_digest(hash_pw(pw, salt)[0], h)


def tok_hash(t):
    return hashlib.sha256(t.encode()).hexdigest()


def level_perms(c, level_id):
    r = c.execute('SELECT perms FROM levels WHERE id=?', (level_id,)).fetchone()
    return json.loads(r['perms']) if r else {'kons': [], 'pdf': False, 'export': False, 'suggest': False}


def user_public(c, u):
    p = level_perms(c, u['level_id'])
    if u['is_admin']:
        p = {'kons': ['*'], 'pdf': True, 'export': True, 'suggest': True}
    lv = c.execute('SELECT name FROM levels WHERE id=?', (u['level_id'],)).fetchone()
    return {'id': u['id'], 'login': u['login'], 'name': u['name'], 'is_admin': bool(u['is_admin']),
            'level': lv['name'] if lv else '', 'perms': p, 'offline_days': u['offline_days']}


def log(c, user, type_, data=None, request=None, ts=None, device=None, offline=0):
    c.execute('INSERT INTO events(ts,recv,user_id,login,type,data,device,ip,offline) VALUES(?,?,?,?,?,?,?,?,?)',
              (ts or now(), now(), user['id'] if user else None, user['login'] if user else (data or {}).get('login'),
               type_, json.dumps(data or {}, ensure_ascii=False), device,
               (request.headers.get('x-forwarded-for') or request.client.host) if request else None, offline))


# ---------- данные ----------
def active_version(c):
    return c.execute('SELECT * FROM versions WHERE active=1 ORDER BY id DESC LIMIT 1').fetchone()


def _hdr(ws):
    return [str(x.value or '').strip().lower() for x in ws[1]]


def _col(H, *names):
    for i, h in enumerate(H):
        if any(h.startswith(n) for n in names):
            return i
    return None


def inspect_zip(raw):
    """Проверяет архив и возвращает (asof, список конструктивов)."""
    z = zipfile.ZipFile(io.BytesIO(raw))
    xl = [n for n in z.namelist() if n.lower().endswith('.xlsx') and not n.split('/')[-1].startswith('~')]
    if not xl:
        raise ValueError('В архиве нет файла Excel')
    wb = openpyxl.load_workbook(io.BytesIO(z.read(xl[0])), read_only=True)
    if 'Требования' not in wb.sheetnames:
        raise ValueError('В Excel нет листа «Требования»')
    ws = wb['Требования']
    rows = list(ws.iter_rows(values_only=True))
    H = [str(x or '').strip().lower() for x in rows[0]]
    ck = _col(H, 'конструктив')
    if ck is None:
        raise ValueError('На листе «Требования» нет столбца «Конструктив»')
    kons = []
    for r in rows[1:]:
        k = str(r[ck] or '').strip()
        if k and k not in kons:
            kons.append(k)
    asof = ''
    if 'Инфо' in wb.sheetnames:
        for r in wb['Инфо'].iter_rows(values_only=True):
            if r and str(r[0] or '').lower().startswith('актуально'):
                v = r[1]
                asof = v.strftime('%d.%m.%Y') if hasattr(v, 'strftime') else str(v or '')
    return asof, kons


_pack_cache = {}


def build_pack(version_id, perms):
    """ZIP только с разрешёнными конструктивами и их документами."""
    key = (version_id, json.dumps(perms, sort_keys=True, ensure_ascii=False))
    if key in _pack_cache:
        return _pack_cache[key]
    raw = open(os.path.join(DATA, 'versions', f'{version_id}.zip'), 'rb').read()
    z = zipfile.ZipFile(io.BytesIO(raw))
    names = z.namelist()
    xname = [n for n in names if n.lower().endswith('.xlsx') and not n.split('/')[-1].startswith('~')][0]
    wb = openpyxl.load_workbook(io.BytesIO(z.read(xname)))
    allow_all = '*' in perms.get('kons', [])
    allowed = set(perms.get('kons', []))
    ws = wb['Требования']
    H = _hdr(ws)
    ck, cd = _col(H, 'конструктив'), _col(H, 'документ', 'источник')
    used = set()
    for i in range(ws.max_row, 1, -1):
        k = str(ws.cell(i, ck + 1).value or '').strip()
        if not k and not any(ws.cell(i, j + 1).value for j in range(len(H))):
            continue
        if allow_all or k in allowed:
            if cd is not None:
                for d in str(ws.cell(i, cd + 1).value or '').split(';'):
                    if d.strip():
                        used.add(d.strip())
        else:
            ws.delete_rows(i)
    files = set()
    if 'Документы' in wb.sheetnames:
        wd = wb['Документы']
        HD = _hdr(wd)
        cc, cf = _col(HD, 'код'), _col(HD, 'файл')
        for i in range(wd.max_row, 1, -1):
            code = str(wd.cell(i, cc + 1).value or '').strip()
            if not code:
                continue
            if code not in used:
                wd.delete_rows(i)
            elif cf is not None and wd.cell(i, cf + 1).value:
                files.add(str(wd.cell(i, cf + 1).value).strip().lower())
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zo:
        b = io.BytesIO(); wb.save(b)
        zo.writestr('Требования_ОКК.xlsx', b.getvalue())
        if perms.get('pdf', True):
            for n in names:
                if n.lower().endswith('.pdf') and n.split('/')[-1].lower() in files:
                    zo.writestr('Документы/' + n.split('/')[-1], z.read(n))
    data = out.getvalue()
    if len(_pack_cache) > 40:
        _pack_cache.clear()
    _pack_cache[key] = data
    return data


def encrypt(dk_b64, data):
    iv = secrets.token_bytes(12)
    return b'OKK2' + iv + AESGCM(base64.b64decode(dk_b64)).encrypt(iv, data, None)


# ---------- приложение ----------
app = FastAPI(title='Требования ОКК', docs_url=None, redoc_url=None, openapi_url=None)
init_db()
_fails = {}


def auth(request: Request):
    h = request.headers.get('authorization', '')
    if not h.startswith('Bearer '):
        raise HTTPException(401, 'Нужен вход')
    with db() as c:
        s = c.execute('SELECT * FROM sessions WHERE token=?', (tok_hash(h[7:]),)).fetchone()
        if not s:
            raise HTTPException(401, 'Сеанс закончился, войдите снова')
        u = c.execute('SELECT * FROM users WHERE id=?', (s['user_id'],)).fetchone()
        if not u or not u['active']:
            raise HTTPException(401, 'Доступ отключён')
        c.execute('UPDATE sessions SET last_seen=? WHERE token=?', (now(), s['token']))
        c.execute('UPDATE users SET last_seen=? WHERE id=?', (now(), u['id']))
        return dict(u)


def admin(u=Depends(auth)):
    if not u['is_admin']:
        raise HTTPException(403, 'Только для администратора')
    return u


@app.post('/api/login')
def login(request: Request, body: dict = Body(...)):
    lg = str(body.get('login', '')).strip().lower()
    pw = str(body.get('password', ''))
    dev = str(body.get('device', ''))[:200]
    ip = request.headers.get('x-forwarded-for') or request.client.host
    k = (lg, ip)
    f = [t for t in _fails.get(k, []) if t > time.time() - 600]
    if len(f) >= 6:
        raise HTTPException(429, 'Слишком много попыток. Подождите 10 минут.')
    with db() as c:
        u = c.execute('SELECT * FROM users WHERE login=?', (lg,)).fetchone()
        if not u or not check_pw(pw, u['pw_hash'], u['pw_salt']):
            _fails[k] = f + [time.time()]
            log(c, None, 'login_fail', {'login': lg}, request, device=dev)
            c.commit()
            raise HTTPException(401, 'Неверный логин или пароль')
        if not u['active']:
            log(c, u, 'login_blocked', {}, request, device=dev)
            c.commit()
            raise HTTPException(401, 'Доступ отключён')
        _fails.pop(k, None)
        t = secrets.token_urlsafe(32)
        c.execute('INSERT INTO sessions(token,user_id,device,created,last_seen) VALUES(?,?,?,?,?)', (tok_hash(t), u['id'], dev, now(), now()))
        c.execute('UPDATE users SET last_seen=? WHERE id=?', (now(), u['id']))
        log(c, u, 'login', {}, request, device=dev)
        v = active_version(c)
        return {'token': t, 'user': user_public(c, u), 'dk': u['dk'], 'version': {'id': v['id'], 'asof': v['asof']} if v else None}


@app.post('/api/check')
def check(u=Depends(auth)):
    with db() as c:
        v = active_version(c)
        return {'user': user_public(c, u), 'version': {'id': v['id'], 'asof': v['asof']} if v else None}


@app.post('/api/logout')
def logout(request: Request, u=Depends(auth)):
    with db() as c:
        c.execute('DELETE FROM sessions WHERE token=?', (tok_hash(request.headers['authorization'][7:]),))
        log(c, u, 'logout', {}, request)
    return {'ok': True}


@app.get('/api/pack')
def pack(request: Request, as_user: int = 0, u=Depends(auth)):
    with db() as c:
        v = active_version(c)
        if not v:
            raise HTTPException(404, 'Данные ещё не загружены')
        target = u
        if as_user and u['is_admin']:
            target = c.execute('SELECT * FROM users WHERE id=?', (as_user,)).fetchone()
            if not target:
                raise HTTPException(404, 'Нет такого пользователя')
            target = dict(target)
        perms = user_public(c, target)['perms']
        log(c, u, 'pack', {'version': v['id'], **({'as': target['login']} if as_user else {})}, request)
    data = build_pack(v['id'], perms)
    return Response(encrypt(u['dk'], data), media_type='application/octet-stream',
                    headers={'X-Version': str(v['id']), 'X-Asof': v['asof'] or '', 'Cache-Control': 'no-store',
                             'X-Perms': base64.b64encode(json.dumps(perms).encode()).decode()})


@app.post('/api/events')
def events(request: Request, body: list = Body(...), u=Depends(auth)):
    with db() as c:
        for e in body[:500]:
            if not isinstance(e, dict):
                continue
            log(c, u, str(e.get('type', 'event'))[:40], e.get('data') if isinstance(e.get('data'), dict) else {}, request,
                ts=str(e.get('ts') or now())[:40], device=str(e.get('device', ''))[:200], offline=1 if e.get('offline') else 0)
    return {'ok': True}


@app.post('/api/suggest')
def suggest(request: Request, body: dict = Body(...), u=Depends(auth)):
    with db() as c:
        if not user_public(c, u)['perms'].get('suggest'):
            raise HTTPException(403, 'Нет права предлагать изменения')
        txt = str(body.get('text', '')).strip()
        if not txt:
            raise HTTPException(400, 'Пустой текст')
        c.execute('INSERT INTO suggestions(ts,user_id,login,kon,card,ids,text) VALUES(?,?,?,?,?,?,?)',
                  (str(body.get('ts') or now())[:40], u['id'], u['login'], str(body.get('kon', ''))[:200], str(body.get('card', ''))[:300],
                   json.dumps(body.get('ids') or [], ensure_ascii=False)[:500], txt[:5000]))
        log(c, u, 'suggest', {'kon': body.get('kon'), 'card': body.get('card')}, request)
    return {'ok': True}


# ---------- администратор ----------
@app.get('/api/admin/state')
def a_state(u=Depends(admin)):
    with db() as c:
        v = active_version(c)
        users = [dict(r) for r in c.execute('SELECT id,login,name,level_id,active,is_admin,offline_days,created,last_seen,note FROM users ORDER BY login')]
        levels = [{'id': r['id'], 'name': r['name'], 'perms': json.loads(r['perms'])} for r in c.execute('SELECT * FROM levels ORDER BY id')]
        versions = [dict(r) for r in c.execute('SELECT id,ts,fname,asof,size,active,by_login FROM versions ORDER BY id DESC LIMIT 20')]
        newsug = c.execute("SELECT COUNT(*) n FROM suggestions WHERE status='new'").fetchone()['n']
        return {'users': users, 'levels': levels, 'kons': json.loads(v['kons']) if v else [], 'versions': versions, 'new_suggestions': newsug}


@app.post('/api/admin/user')
def a_user(request: Request, body: dict = Body(...), u=Depends(admin)):
    lg = str(body.get('login', '')).strip().lower()
    with db() as c:
        if body.get('id'):
            old = c.execute('SELECT * FROM users WHERE id=?', (body['id'],)).fetchone()
            if not old:
                raise HTTPException(404, 'Нет такого пользователя')
            if old['id'] == u['id'] and (not body.get('active', True) or not body.get('is_admin', True)):
                raise HTTPException(400, 'Нельзя отключить или лишить прав самого себя')
            c.execute('UPDATE users SET name=?,level_id=?,active=?,is_admin=?,offline_days=?,note=? WHERE id=?',
                      (body.get('name', ''), body.get('level_id'), 1 if body.get('active', True) else 0, 1 if body.get('is_admin') else 0,
                       int(body.get('offline_days') or 35), body.get('note', ''), body['id']))
            if not body.get('active', True):
                c.execute('DELETE FROM sessions WHERE user_id=?', (body['id'],))
            if body.get('password'):
                h, s = hash_pw(body['password'])
                c.execute('UPDATE users SET pw_hash=?,pw_salt=? WHERE id=?', (h, s, body['id']))
                c.execute('DELETE FROM sessions WHERE user_id=?', (body['id'],))
            log(c, u, 'admin_user_edit', {'user': old['login'], 'password_changed': bool(body.get('password'))}, request)
            return {'ok': True}
        if not lg or not body.get('password'):
            raise HTTPException(400, 'Нужны логин и пароль')
        if c.execute('SELECT 1 FROM users WHERE login=?', (lg,)).fetchone():
            raise HTTPException(400, 'Такой логин уже есть')
        h, s = hash_pw(body['password'])
        c.execute('INSERT INTO users(login,name,pw_hash,pw_salt,level_id,active,is_admin,offline_days,dk,created,note) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                  (lg, body.get('name', ''), h, s, body.get('level_id'), 1, 1 if body.get('is_admin') else 0, int(body.get('offline_days') or 35),
                   base64.b64encode(secrets.token_bytes(32)).decode(), now(), body.get('note', '')))
        log(c, u, 'admin_user_add', {'user': lg}, request)
    return {'ok': True}


@app.post('/api/admin/level')
def a_level(request: Request, body: dict = Body(...), u=Depends(admin)):
    p = body.get('perms') or {}
    perms = {'kons': list(p.get('kons') or []), 'pdf': bool(p.get('pdf')), 'export': bool(p.get('export')), 'suggest': bool(p.get('suggest'))}
    with db() as c:
        if body.get('delete'):
            if c.execute('SELECT 1 FROM users WHERE level_id=?', (body['id'],)).fetchone():
                raise HTTPException(400, 'Уровень назначен пользователям — сначала смените им уровень')
            c.execute('DELETE FROM levels WHERE id=?', (body['id'],))
            log(c, u, 'admin_level_delete', {'id': body['id']}, request)
            return {'ok': True}
        name = str(body.get('name', '')).strip()
        if not name:
            raise HTTPException(400, 'Нужно название уровня')
        if body.get('id'):
            c.execute('UPDATE levels SET name=?,perms=? WHERE id=?', (name, json.dumps(perms, ensure_ascii=False), body['id']))
        else:
            c.execute('INSERT INTO levels(name,perms) VALUES(?,?)', (name, json.dumps(perms, ensure_ascii=False)))
        log(c, u, 'admin_level_save', {'name': name, 'perms': perms}, request)
    return {'ok': True}


def _events_q(c, user_id, d_from, d_to, type_, q, limit):
    sql, a = 'SELECT * FROM events WHERE 1=1', []
    if user_id:
        sql += ' AND user_id=?'; a.append(user_id)
    if d_from:
        sql += ' AND ts>=?'; a.append(d_from)
    if d_to:
        sql += ' AND ts<?'; a.append(d_to + 'T99')
    if type_:
        sql += ' AND type=?'; a.append(type_)
    if q:
        sql += ' AND (data LIKE ? OR login LIKE ?)'; a += ['%' + q + '%', '%' + q + '%']
    sql += ' ORDER BY ts DESC, id DESC LIMIT ?'; a.append(limit)
    return [dict(r) for r in c.execute(sql, a)]


@app.get('/api/admin/events')
def a_events(user_id: int = 0, d_from: str = '', d_to: str = '', type: str = '', q: str = '', limit: int = 500, u=Depends(admin)):
    with db() as c:
        return _events_q(c, user_id, d_from, d_to, type, q, min(limit, 5000))


@app.get('/api/admin/events.csv')
def a_events_csv(user_id: int = 0, d_from: str = '', d_to: str = '', type: str = '', q: str = '', u=Depends(admin)):
    with db() as c:
        rows = _events_q(c, user_id, d_from, d_to, type, q, 100000)
    b = io.StringIO()
    w = csv.writer(b, delimiter=';')
    w.writerow(['Время', 'Получено', 'Логин', 'Действие', 'Подробности', 'Устройство', 'IP', 'Офлайн'])
    for r in rows:
        w.writerow([r['ts'], r['recv'], r['login'], r['type'], r['data'], r['device'], r['ip'], 'да' if r['offline'] else ''])
    return Response('﻿' + b.getvalue(), media_type='text/csv; charset=utf-8',
                    headers={'Content-Disposition': 'attachment; filename="okk-journal.csv"'})


@app.get('/api/admin/suggestions')
def a_sugs(u=Depends(admin)):
    with db() as c:
        return [dict(r) for r in c.execute('SELECT * FROM suggestions ORDER BY id DESC LIMIT 500')]


@app.post('/api/admin/suggestion')
def a_sug(request: Request, body: dict = Body(...), u=Depends(admin)):
    st = body.get('status')
    if st not in ('new', 'accepted', 'rejected'):
        raise HTTPException(400, 'Неверный статус')
    with db() as c:
        c.execute('UPDATE suggestions SET status=?,note=?,decided=? WHERE id=?', (st, body.get('note', ''), now(), body['id']))
        log(c, u, 'admin_suggestion', {'id': body['id'], 'status': st}, request)
    return {'ok': True}


@app.post('/api/admin/upload')
async def a_upload(request: Request, file: UploadFile = File(...), u=Depends(admin)):
    raw = await file.read()
    try:
        asof, kons = inspect_zip(raw)
    except Exception as e:
        raise HTTPException(400, f'Архив не принят: {e}')
    with _lock, db() as c:
        cur = c.execute('INSERT INTO versions(ts,fname,asof,size,kons,active,by_login) VALUES(?,?,?,?,?,0,?)',
                        (now(), file.filename, asof, len(raw), json.dumps(kons, ensure_ascii=False), u['login']))
        vid = cur.lastrowid
        open(os.path.join(DATA, 'versions', f'{vid}.zip'), 'wb').write(raw)
        c.execute('UPDATE versions SET active=CASE WHEN id=? THEN 1 ELSE 0 END', (vid,))
        log(c, u, 'admin_upload', {'version': vid, 'file': file.filename, 'asof': asof}, request)
    _pack_cache.clear()
    return {'ok': True, 'version': vid, 'asof': asof, 'kons': kons}


@app.post('/api/admin/activate')
def a_activate(request: Request, body: dict = Body(...), u=Depends(admin)):
    with db() as c:
        if not c.execute('SELECT 1 FROM versions WHERE id=?', (body['id'],)).fetchone():
            raise HTTPException(404, 'Нет такой версии')
        c.execute('UPDATE versions SET active=CASE WHEN id=? THEN 1 ELSE 0 END', (body['id'],))
        log(c, u, 'admin_activate', {'version': body['id']}, request)
    _pack_cache.clear()
    return {'ok': True}


@app.get('/api/health')
def health():
    return {'ok': True}


@app.get('/admin')
def admin_page():
    return FileResponse(os.path.join(WEB, 'admin.html'), headers={'Cache-Control': 'no-cache'})


@app.middleware('http')
async def no_cache_html(request: Request, call_next):
    r = await call_next(request)
    p = request.url.path
    if p == '/' or p.endswith('.html') or p.endswith('sw.js') or p.endswith('.webmanifest'):
        r.headers['Cache-Control'] = 'no-cache'
    r.headers['X-Content-Type-Options'] = 'nosniff'
    r.headers['Referrer-Policy'] = 'same-origin'
    return r


if os.path.isdir(WEB):
    app.mount('/', StaticFiles(directory=WEB, html=True), name='web')
